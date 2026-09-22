"""Populate copied course templates; never modifies the original appendix files.

The defect input must first be a Word-converted copy of appendix 2 (.doc -> .docx).
This generator writes the two final deliverables in module2-deliverables/.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "module2-deliverables"
DEFECT = OUT / "PEFT_模块二_缺陷报告.docx"
REPORT = OUT / "PEFT_模块二_测试报告.docx"
REPORT_TEMPLATE = ROOT.parent / "实践作业-文档模板2026" / "附录4：测试报告模板（模块二）.docx"


def put(paragraph, content):
    paragraph.text = content
    paragraph.paragraph_format.space_after = Pt(6)


def cell(table, row, column, content):
    table.cell(row, column).text = content


def revision(table):
    for col, content in enumerate(("2026-09-20", "V1.0", "完成模块二测试、缺陷修复与验证", "Djie1234 / AI辅助")):
        cell(table, 1, col, content)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def build_report():
    doc = Document(REPORT_TEMPLATE)
    p = doc.paragraphs
    put(p[3], "PEFT 任务调度算法")
    put(p[5], "模块二 AI 辅助测试报告")
    revision(doc.tables[0])
    put(p[26], "被测对象为 PEFT（Predict Earliest Finish Time）有向无环任务图调度项目。系统以任务 DAG、计算时间矩阵、处理器通信带宽矩阵及已有任务队列为输入，返回每个处理器的排程、任务事件和前驱映射。本次以已上传模块一基线 d38620b 为起点，重点测试 peft/peft.py 中 schedule_dag 与 _compute_eft 的已有任务空隙插入逻辑，同时检查输入契约和输出副作用。")
    put(p[27], "选择课程方案 2“AI 测”：使用 Codex AI 助手提出候选场景、编写 pytest 代码、比较实际与预期、定位缺陷并生成文档。AI 不作为正确性判据；预期结果来自调度函数“最早可插入位置”的定义、无重叠约束和输入校验契约，并由可重复运行的断言验证。")
    put(p[28], "AI 使用范围：场景构思、测试代码草拟、失败解释、最小修复建议及文档起草。被测系统未嵌入 AI 模块，故本项不适用。")
    put(p[29], "工具为 Codex AI 编程助手（服务端模型版本可能随平台配置变化，不把未知版本写成固定值）。输入为本项目代码、模块一 36 条现有用例及作业要求；输出为候选清单、20 条新用例和同号脚本、失败/回归记录。关键过程保存在 AI过程记录.md。")
    put(p[32], "实施过程：① 从 d38620b 建独立工作树；② AI 提出 23 个候选，排除 3 个与模块一 TC-009、025、026 重复的场景，保留 20 条；③ 在未修复代码上执行，20 条中 16 通过、4 失败；④ 对同一根因进行最小修复；⑤ 再执行模块二 20 条及模块一 36 条、原项目 2 条，总计 58 条通过。测试设计混用边界值、无效等价类、状态保持和回归测试，脚本与 Excel 编号一一对应。")
    put(p[33], "AI 原始建议与最终成果的差异：原始 23 个候选含 3 个模块一重复场景；最终去重后为 20 条。四个失败经代码检查归并为一个首空隙 exact-fit 缺陷。原始建议、删改理由及命令在 AI过程记录.md，原始 JUnit 见 evidence/baseline-fail.xml。颜子轩和余璨共同审阅最终用例与缺陷结论。")
    put(p[35], "新增用例 M2-TC-001～020 均在附录 1 格式的 PEFT_模块二_测试用例清单.xlsx 中列出，以下为编号、意图及修复后结果。用例函数位于 tests_module2/test_module2.py；全部 20/20 通过。")
    put(p[36], "确认缺陷 DEF-M2-001：已有事件 [2,4] 时，耗时 2 的新任务应插入 [0,2]，基线却排在 [4,6]。代码首空隙判断采用 > 0，错误排除了恰好容纳的边界；内部空隙已采用 >= 0。修复仅将首空隙条件改为 >= 0。修复前 M2-TC-001、002、003、009 失败，修复后四条通过，全部 58 条回归通过；未观察到第二个可确证缺陷。详见缺陷报告和三份 JUnit XML。")
    put(p[38], "执行环境：Windows 11 家庭版中文版（Build 26200，64 位）；Python 3.11.2；pytest 7.1.2；NumPy 1.26.4；NetworkX 2.8.5；Matplotlib 3.7.5。硬件内存约 15.4 GB。项目使用本地解释器执行，无数据库、外部服务或网络连接依赖。")
    put(p[39], "运行命令：在项目根目录执行 python run_module2_tests.py。该入口调用 pytest 运行 tests_module2、tests_module1 与 peft/test，输出 JUnit 与 HTML 覆盖率。实测 58 passed，整体语句覆盖率 82%；原始 JUnit 文件在 module2-deliverables/evidence/，可独立复核。")
    put(p[41], "反思一：AI 系统在测试设计中的不足。被测 PEFT 是确定性算法，本轮讨论的 AI 是测试辅助工具而非被测算法；因此无法用 58 条测试证明 AI 输出永远正确，只能证明这些输入下的具体断言。")
    put(p[42], "AI 的有效处：能快速扫描已有测试与代码分支，提出 exact-fit、非零偏移、矩阵维度等候选，并生成 20 个可执行函数，使边界缺陷在修复前稳定重现。")
    put(p[43], "AI 的错误与局限：候选中 3 条与模块一重复；同一缺陷的 4 个失败容易被误计为 4 个缺陷。若只接受模型的自然语言判断而不先运行基线，就可能将猜测写成结果。")
    put(p[44], "测试挑战：调度预期需同时考虑任务准备时间、通信延迟和处理器空隙，复杂 DAG 很难人工给出唯一完整排程。本轮采用最小单处理器反例构造明确 oracle，并用全量回归防止副作用。")
    put(p[45], "结论：AI 可提高候选生成和脚本编写速度，但输出要经排重、可计算预期和红绿回归核对；不能仅以 AI 自信程度替代验证，也不能从有限测试推导整个系统无缺陷。")
    put(p[47], "反思二：AI 作为测试工具的效用与风险。")
    put(p[48], "收益：辅助识别 > 与 >= 的边界差异，快速把数值、偏移和映射结果扩展为可重复的测试，并生成执行与报告草稿。")
    put(p[49], "陷阱：可能重复已有用例，或把同一根因夸大成多个缺陷。为此保留原始候选、pytest XML、环境版本与可复现命令，并通过小组审阅核查关键结论。")
    put(p[50], "关键修正：剔除 3 条重复建议；将 4 条失败合并为 DEF-M2-001；最小代码改动为首空隙判定 >= 0；修复后重新执行全部 58 条测试。")
    put(p[51], "角色定位：AI 是候选生成与实现助手，测试结论由小组根据输入、代码与执行证据复核。最终结论以可重复的自动化结果和实际排程为依据。")
    put(p[53], "颜子轩与余璨合作完成模块二，贡献各占 50%。颜子轩侧重自动化测试、最小修复、运行记录和证据整理；余璨侧重测试用例复核、缺陷分析与报告整理。双方共同审阅最终测试结论。")

    titles = [
        "首空隙整时 exact-fit", "小数 exact-fit", "非零偏移 exact-fit", "首空隙较大", "首空隙较小",
        "内部空隙 exact-fit", "内部空隙偏移", "已有事件保留", "前驱映射", "DAG 无副作用",
        "非方阵通信矩阵", "通信 Inf", "负通信值", "零处理器", "多根 DAG", "多终点 DAG",
        "禁重标号非连续节点", "无向图", "一维计算矩阵", "元组矩阵正规化",
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c, v in zip(table.rows[0].cells, ("用例编号", "测试意图", "修复后")):
        c.text = v
    set_repeat_header(table.rows[0])
    for i, title in enumerate(titles, 1):
        cells = table.add_row().cells
        for c, v in zip(cells, (f"M2-TC-{i:03d}", title, "通过")):
            c.text = v
    table.columns[0].width = Cm(3)
    table.columns[1].width = Cm(9)
    table.columns[2].width = Cm(2.4)
    p[37]._p.addprevious(table._tbl)
    doc.save(REPORT)


def build_defect():
    doc = Document(DEFECT)
    p = doc.paragraphs
    put(p[3], "PEFT 任务调度算法")
    put(p[5], "模块二缺陷报告：首空隙 exact-fit 被遗漏")
    revision(doc.tables[0])
    refs = doc.tables[1]
    for col, value in enumerate(("1", "基线代码", "d38620b", "模块一已上传基线", "Git 历史")):
        cell(refs, 1, col, value)
    for col, value in enumerate(("2", "新增测试", "tests_module2/test_module2.py", "M2-TC-001/002/003/009", "本分支")):
        cell(refs, 2, col, value)
    put(p[36], "记录基线 d38620b 在模块二边界测试中发现的可稳定复现缺陷，给出输入、预期、实际、原因、修复及回归证据。受众为课程教师、测试小组与后续维护者。")
    put(p[37], "项目名称：PEFT 模块二 AI 辅助测试。")
    put(p[38], "测试记录：M2-TC-001、002、003、009；自动化文件 tests_module2/test_module2.py。")
    put(p[39], "测试由项目组在本地工作树中执行；颜子轩与余璨共同审阅用例与缺陷结论。")
    put(p[41], "PEFT 为 Python/NetworkX DAG 任务调度项目。本次聚焦 peft/peft.py::_compute_eft 对已有处理器任务队列的最早空隙插入。基线版本 d38620b；修复版本见本分支下一次 fix 提交。")
    put(p[42], "项目/产品名称：PEFT。")
    put(p[43], "测试环境：Windows 11 家庭版中文版 Build 26200，Python 3.11.2，pytest 7.1.2，NumPy 1.26.4，NetworkX 2.8.5。")
    put(p[44], "数据库：不适用。")
    put(p[45], "用户：本地命令行测试执行者。")
    put(p[46], "网络：不依赖外部服务。")
    put(p[48], "exact-fit：空隙长度恰好等于新任务执行时长；EFT：最早完成时间；JUnit XML：pytest 生成的原始执行结果。")
    put(p[50], "PEFT 项目代码、模块一测试、模块二 20 条新增测试及 evidence 下三份 pytest JUnit XML。")
    put(p[53], "64 位 Windows 11；内存约 15.4 GB；CPU 型号不影响本缺陷复现。")
    put(p[55], "Python 3.11.2、pytest 7.1.2；本地解释器运行，测试不依赖数据库或浏览器。")
    put(p[56], "3 目标场景测试与缺陷验证")
    put(p[58], "测试版本：基线 d38620b；修复版本为本模块二分支。所有文件相对项目根目录。")
    put(p[60], "选取首空隙 exact-fit 作为关键边界，配以小数时长、非零偏移及字典映射对照；同时执行模块二其余 16 条及模块一回归，验证修复没有破坏既有行为。")
    put(p[62], "基线运行 python -m pytest tests_module2 -q：16 通过、4 失败，见 evidence/baseline-fail.xml。仅将 peft/peft.py 首空隙条件由 > 0 改为 >= 0；再运行模块二 20/20 通过（fixed-module2.xml），全量 58/58 通过（fixed-all.xml）。")
    put(p[64], "")
    p[64].style = doc.styles["Normal"]
    put(p[65], "以下记录对应一个确认缺陷，不将四个关联用例分别计为四个缺陷。")
    put(p[66], "3.4.1  DEF-M2-001 处理器首空隙 exact-fit 被遗漏")
    put(p[67], "缺陷数 1；修复前模块二 20 条中 4 条失败，均指向同一根因；修复后模块二 20/20、全量 58/58 通过。当前状态为已修复并经自动化回归验证，小组已共同审阅缺陷结论。")
    put(p[68], "建议后续对多处理器、大型 DAG 和浮点容差继续扩充测试；本报告不推断所有输入均正确。")
    t = doc.tables[2]
    cell(t, 0, 1, "项目组自动化执行")
    cell(t, 0, 7, "2026-09-20")
    cell(t, 1, 1, "_compute_eft 空隙插入")
    cell(t, 1, 7, "PEFT-SCHED-01")
    cell(t, 2, 1, "调度最早空隙")
    cell(t, 3, 1, "任务可在处理器首空隙内无重叠运行时，应使用该最早可行位置。")
    cell(t, 4, 1, "M2-TC-001/002/003/009")
    cell(t, 5, 1, "中")
    cell(t, 5, 6, "中")
    cell(t, 5, 10, "已修复/自动回归通过")
    cell(t, 6, 1, "项目组最小修复")
    cell(t, 7, 1, "颜子轩、余璨")
    cell(t, 8, 1, "首空隙长度恰等于任务时长时，任务错误推迟到已有事件之后")
    cell(t, 9, 0, "复现步骤：1. 单处理器已有 ScheduleEvent(0,2,4,0)。2. 对单节点 DAG 调用 schedule_dag，计算矩阵 [[2]]、通信矩阵 [[0]]、time_offset=0。\n预期结果：新任务 1 在 [0,2]，已有任务仍在 [2,4]。\n基线实际：新任务在 [4,6]，错过完全可用的 [0,2]。\n扩展验证：小数 exact-fit 与非零偏移场景同样失败；字典映射随排程错误。")
    cell(t, 10, 2, "evidence/baseline-fail.xml；evidence/fixed-module2.xml；evidence/fixed-all.xml")
    cell(t, 11, 2, "无其他已确认关联缺陷")
    cell(t, 12, 2, "四个失败用例归于同一根因。")
    cell(t, 14, 2, "项目组最小修复")
    cell(t, 14, 8, "2026-09-20")
    cell(t, 15, 2, "模块二 fix 提交")
    cell(t, 15, 8, "最小边界修复")
    cell(t, 16, 2, "peft/peft.py::_compute_eft 在检查首个已有任务前的可插入空隙时，原条件 (start - duration) - ready_time > 0 排除等于 0 的可行情况；改为 >= 0，和内部空隙规则一致。20 条新增测试、58 条全量测试通过。")
    cell(t, 17, 2, "自动化验证通过；小组审阅确认")
    doc.save(DEFECT)


if __name__ == "__main__":
    build_report()
    build_defect()
    print(REPORT)
    print(DEFECT)
